"""XLSX and JSON through the same request, job, limits and audit as the CSV:
the format is a parameter, never a path of its own."""

import io
import json
import zipfile

import pytest
from common.models.user import User
from common.records import RecordId
from httpx import AsyncClient
from openpyxl import load_workbook

from backend.app.container import container
from backend.app.models.enum.workspace_roles import WorkspaceRoles
from backend.app.models.filter_queries.response_export import ResponseExportFilters
from backend.app.schemas.response_export import ExportKind, ExportStatus
from backend.app.schemas.workspace_user import WorkspaceUserDocument
from backend.app.services.response_export import limits
from backend.app.services.response_export import xlsx as xlsx_module
from tests.app.controllers.data import invited_user, testUser
from tests.app.controllers.test_permission_matrix import _cookies
from tests.app.registration.conftest import memory_jobs  # noqa: F401
from tests.app.response_export.helpers import (
    GROUP_FIELDS,
    answers,
    publish,
    rows,
    start,
    submit,
    submit_many,
    url,
)
from tests.app.response_export.test_export_job import build, fetch, queued

pytestmark = pytest.mark.asyncio

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def audit(workspace, form):
    return container.response_export_repo().list_for_form(workspace.id, form.form_id)


def sheet_rows(content, name: str = "Responses"):
    content = getattr(content, "content", content)
    return [list(r) for r in load_workbook(io.BytesIO(content))[name].iter_rows(values_only=True)]


@pytest.fixture
def small_threshold(monkeypatch, memory_jobs):  # noqa: F811
    """Two responses already make a job; a chunk holds a few lines."""
    monkeypatch.setattr(limits, "STREAM_MAX_ROWS", 2)
    monkeypatch.setattr(limits, "CHUNK_BYTES", 400)
    return memory_jobs


# ------------------------------------------------------------------ streaming


async def test_a_small_export_streams_an_excel_workbook(client: AsyncClient, workspace):
    form = await publish(workspace, hidden=["utm_source"])
    await submit(workspace, form, answers("Ada"), hidden={"utm_source": "mail"})
    await submit(workspace, form, answers("=SUM(1)", phone="+31 6 99"))
    response = await start(client, workspace, form, {"format": "xlsx"})
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"
    assert 'filename="Export test.xlsx"' in response.headers["content-disposition"]
    table = sheet_rows(response.content)
    assert table[0] == [
        "Responder ID", "Full name", "Phone", "Age", "Colours",
        "Response ID", "Submitted at (UTC, ISO 8601)", "Hidden: utm_source",
    ]
    by_name = {r[1]: r for r in table[1:]}
    assert by_name["Ada"][0] == testUser.sub and by_name["Ada"][2:5] == ["+31 6 1234", 30, "Red; Blue"]
    # the text is as written (no apostrophe): it is a string cell, never a formula
    assert by_name["=SUM(1)"][2] == "+31 6 99"
    assert by_name["Ada"][6].year >= 2026 and by_name["Ada"][7] == "mail"
    (entry,) = await audit(workspace, form)
    assert entry.format == "xlsx" and entry.row_count == 2 and entry.kind == ExportKind.STREAM


async def test_a_small_export_streams_json(client: AsyncClient, workspace):
    form = await publish(workspace, hidden=["utm_source"])
    await submit(workspace, form, answers("Ada"), hidden={"utm_source": "mail"})
    response = await start(client, workspace, form, {"format": "json"})
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/json")
    assert 'filename="Export test.json"' in response.headers["content-disposition"]
    assert not response.content.startswith(b"\xef\xbb\xbf")
    document = json.loads(response.content)
    assert document["form"]["id"] == form.form_id and document["form"]["title"] == "Export test"
    (item,) = document["responses"]
    assert item["responder"] == testUser.sub and item["formVersion"] in (None, 1)
    assert item["answers"]["name"]["text"] == "Ada"
    assert item["answers"]["colours"]["choices"]["values"] == ["c1", "c3"]
    assert item["hiddenFields"] == {"utm_source": "mail"}
    assert [q["id"] for q in document["questions"]] == ["name", "phone", "age", "colours"]
    (entry,) = await audit(workspace, form)
    assert entry.format == "json" and entry.row_count == 1
    assert "Ada" not in json.dumps(entry.model_dump(mode="json"))


async def test_the_default_format_is_still_csv(client: AsyncClient, workspace):
    form = await publish(workspace)
    await submit(workspace, form)
    for body in ({}, {"format": "csv"}):
        response = await start(client, workspace, form, body)
        assert response.headers["content-type"].startswith("text/csv")
        assert rows(response)[0][0] == "Responder ID"


async def test_every_format_takes_the_same_columns_and_filters(client: AsyncClient, workspace):
    form = await publish(workspace)
    await submit_many(workspace, form, 2)
    body = {"columns": ["q:name", "response_id"], "respondent": "known"}
    results = {fmt: await start(client, workspace, form, {**body, "format": fmt}) for fmt in ("csv", "xlsx", "json")}
    assert all(r.status_code == 200 for r in results.values())
    # the filter holds nothing back from (or lets anything into) any of them
    assert len(rows(results["csv"])) == 3 and len(sheet_rows(results["xlsx"])) == 3
    assert len(json.loads(results["json"].content)["responses"]) == 2
    assert [q["id"] for q in json.loads(results["json"].content)["questions"]] == ["name"]
    assert sheet_rows(results["xlsx"])[0] == rows(results["csv"])[0] == ["Full name", "Response ID"]
    entries = await audit(workspace, form)
    assert {e.format for e in entries} == {"csv", "xlsx", "json"}
    assert {tuple(e.columns) for e in entries} == {("q:name", "response_id")}


async def test_an_unknown_format_or_column_is_refused_in_any_format(client: AsyncClient, workspace):
    form = await publish(workspace)
    await submit(workspace, form)
    refused = await start(client, workspace, form, {"format": "pdf"})
    assert refused.status_code == 422
    refused = await start(client, workspace, form, {"format": "xlsx", "columns": ["q:nope"]})
    assert refused.status_code == 422 and refused.json()["code"] == "unknown_column"
    refused = await start(client, workspace, form, {"format": "json", "columns": []})
    assert refused.status_code == 422 and refused.json()["code"] == "no_columns"
    assert await audit(workspace, form) == []
    # the service refuses an unknown format by itself too
    service = container.response_export_service()
    with pytest.raises(Exception) as caught:
        await service.request(workspace.id, form.form_id, testUser, None, ResponseExportFilters(), "pdf")
    assert caught.value.status_code == 422 and caught.value.content["code"] == "unknown_format"


async def test_a_group_form_comes_as_sheets_in_a_workbook_and_nested_in_json(client: AsyncClient, workspace):
    form = await publish(workspace, fields=GROUP_FIELDS)
    given = answers("Ada")
    given["jobs"] = {
        "type": "group",
        "items": [
            {"employer": {"type": "text", "text": "Acme"}},
            {"employer": {"type": "text", "text": "=cmd"}},
        ],
    }
    await submit(workspace, form, given)
    workbook = await start(client, workspace, form, {"format": "xlsx"})
    assert workbook.headers["content-type"] == XLSX
    names = load_workbook(io.BytesIO(workbook.content)).sheetnames
    assert names == ["Responses", "Jobs"]
    jobs = sheet_rows(workbook.content, "Jobs")
    assert jobs[0] == ["Response ID", "Job", "Employer"]
    assert [r[1:] for r in jobs[1:]] == [[1, "Acme"], [2, "=cmd"]]
    main = sheet_rows(workbook.content)
    assert jobs[1][0] == main[1][main[0].index("Response ID")]
    # the audit says what the file is, and how many responses (not group rows) it holds
    (entry,) = await audit(workspace, form)
    assert entry.format == "xlsx" and entry.row_count == 1

    nested = json.loads((await start(client, workspace, form, {"format": "json"})).content)
    items = nested["responses"][0]["answers"]["jobs"]["items"]
    assert [i["employer"]["text"] for i in items] == ["Acme", "=cmd"]


# ------------------------------------------------------------- shared limits


async def test_the_rate_limit_counts_every_format_together(client, workspace, monkeypatch):
    monkeypatch.setattr(limits, "REQUESTS_PER_WINDOW", 3)
    form = await publish(workspace)
    await submit(workspace, form)
    for fmt in ("csv", "xlsx", "json"):
        assert (await start(client, workspace, form, {"format": fmt})).status_code == 200
    refused = await start(client, workspace, form, {"format": "xlsx"})
    assert refused.status_code == 429 and refused.json()["code"] == "rate_limited"
    assert int(refused.headers["retry-after"]) >= 1
    assert len(await audit(workspace, form)) == 3


async def test_too_many_responses_are_refused_in_every_format(client, workspace, monkeypatch):
    monkeypatch.setattr(limits, "MAX_ROWS", 2)
    form = await publish(workspace)
    await submit_many(workspace, form, 3)
    for fmt in ("csv", "xlsx", "json"):
        response = await start(client, workspace, form, {"format": fmt})
        assert response.status_code == 413 and response.json()["code"] == "too_many_rows"
    assert await audit(workspace, form) == []


@pytest.mark.parametrize(
    "role,allowed",
    [
        (WorkspaceRoles.ADMIN, True),
        (WorkspaceRoles.COLLABORATOR, True),
        (WorkspaceRoles.REVIEWER, False),
        (WorkspaceRoles.VIEWER, False),
        (WorkspaceRoles.PRIVACY_OFFICER, False),
    ],
)
async def test_the_permission_is_the_exports_whatever_the_format(client, workspace, role, allowed):
    other = User(id=str(RecordId()), sub="role@example.com")
    await container.workspace_user_repo().save(
        WorkspaceUserDocument(workspace_id=workspace.id, user_id=other.id, roles=[role])
    )
    form = await publish(workspace)
    await submit(workspace, form)
    for fmt in ("xlsx", "json"):
        response = await start(client, workspace, form, {"format": fmt}, user=other)
        assert response.status_code == (200 if allowed else 403), (fmt, response.text)
        assert allowed or "Ada" not in response.text
    assert len(await audit(workspace, form)) == (2 if allowed else 0)


async def test_a_stream_of_any_format_holds_a_slot_while_it_runs(workspace, monkeypatch):
    monkeypatch.setattr(limits, "STREAMS_PER_USER", 1)
    form = await publish(workspace)
    await submit_many(workspace, form, 2)
    service = container.response_export_service()
    first = await service.request(workspace.id, form.form_id, testUser, None, ResponseExportFilters(), "json")
    assert await first.body.__anext__()
    with pytest.raises(Exception) as busy:
        await service.request(workspace.id, form.form_id, testUser, None, ResponseExportFilters(), "xlsx")
    assert busy.value.status_code == 429 and busy.value.content["code"] == "export_busy"
    await first.body.aclose()


async def test_too_many_columns_for_a_sheet_is_refused_for_xlsx_only(client, workspace, monkeypatch):
    monkeypatch.setattr(xlsx_module, "MAX_COLUMNS", 3)
    form = await publish(workspace)
    await submit(workspace, form)
    refused = await start(client, workspace, form, {"format": "xlsx"})
    assert refused.status_code == 422 and refused.json()["code"] == "too_many_columns"
    assert await audit(workspace, form) == []
    assert (await start(client, workspace, form, {"format": "json"})).status_code == 200
    assert (await start(client, workspace, form, {"format": "csv"})).status_code == 200


# ------------------------------------------------------------------------ jobs


async def test_a_large_workbook_is_a_job_with_the_same_file_as_the_stream(client: AsyncClient, workspace, small_threshold, monkeypatch):
    form = await publish(workspace)
    await submit_many(workspace, form, 5)
    export = await queued(client, workspace, form, {"format": "xlsx"})
    assert export["format"] == "xlsx" and export["fileName"] == "Export test.xlsx"
    assert export["status"] == "queued" and export["kind"] == "job"
    assert await build(export["id"]) == "ready"
    file = await fetch(client, workspace, form, export["id"])
    assert file.status_code == 200 and file.headers["content-type"] == XLSX
    assert 'filename="Export test.xlsx"' in file.headers["content-disposition"]
    table = sheet_rows(file.content)
    assert len(table) == 6 and {r[1] for r in table[1:]} == {f"Person {n}" for n in range(5)}
    entry = await container.response_export_repo().get_by_id(export["id"])
    assert entry.format == "xlsx" and entry.row_count == 5 and entry.params_sealed is None
    assert "Person" not in json.dumps(entry.model_dump(mode="json"))
    # many chunks, one valid workbook
    assert entry.chunk_count > 1
    # the streamed file holds the same table
    monkeypatch.setattr(limits, "STREAM_MAX_ROWS", 100)
    streamed = await start(client, workspace, form, {"format": "xlsx"})
    assert [r[:6] for r in sheet_rows(streamed.content)] == [r[:6] for r in table]


async def test_a_large_json_export_is_a_job_too(client: AsyncClient, workspace, small_threshold):
    form = await publish(workspace)
    await submit_many(workspace, form, 4)
    export = await queued(client, workspace, form, {"format": "json", "columns": ["q:name"]})
    assert export["format"] == "json"
    (job,) = small_threshold.jobs.values()
    # only the id travels in the job's arguments: never the format, the columns or the filters
    assert job["args"] == {"export_id": export["id"]}
    assert await build(export["id"]) == "ready"
    file = await fetch(client, workspace, form, export["id"])
    assert file.headers["content-type"].startswith("application/json")
    document = json.loads(file.content)
    assert len(document["responses"]) == 4 and [q["id"] for q in document["questions"]] == ["name"]
    assert all(set(r["answers"]) == {"name"} for r in document["responses"])


async def test_the_format_waits_sealed_with_the_filters(client: AsyncClient, workspace, small_threshold):
    form = await publish(workspace)
    await submit_many(workspace, form, 3)
    export = await queued(client, workspace, form, {"format": "xlsx", "search": testUser.sub})
    stored = await container.response_export_repo().get_by_id(export["id"])
    assert stored.params_sealed and "xlsx" not in stored.params_sealed
    assert stored.format == "xlsx" and stored.filters == {"search": True}
    params = container.response_export_service()._unseal(stored)
    assert params["format"] == "xlsx" and params["filters"]["search"] == testUser.sub


async def test_a_job_queued_before_formats_existed_builds_a_csv(client: AsyncClient, workspace, small_threshold):
    form = await publish(workspace)
    await submit_many(workspace, form, 3)
    export = await queued(client, workspace, form)
    service = container.response_export_service()
    row = await container.response_export_repo().get_by_id(export["id"])
    params = service._unseal(row)
    params.pop("format")
    row.params_sealed = service._seal(row, params)
    await container.response_export_repo().save(row)
    assert await build(export["id"]) == "ready"
    file = await fetch(client, workspace, form, export["id"])
    assert file.headers["content-type"].startswith("text/csv") and rows(file)[0][0] == "Responder ID"


async def test_a_workbook_job_whose_sheet_would_overflow_fails_with_a_code(client: AsyncClient, workspace, small_threshold, monkeypatch):
    form = await publish(workspace, fields=GROUP_FIELDS)
    for number in range(3):
        given = answers(f"P{number}")
        given["jobs"] = {"type": "group", "items": [{"employer": {"type": "text", "text": "E"}}] * 2}
        await submit(workspace, form, given)
    export = await queued(client, workspace, form, {"format": "xlsx"})
    monkeypatch.setattr(xlsx_module, "MAX_ROWS_PER_SHEET", 4)
    assert await build(export["id"]) == "failed"
    entry = await container.response_export_repo().get_by_id(export["id"])
    assert entry.error_code == "too_large" and await container.response_export_repo().count_chunks(export["id"]) == 0


async def test_a_job_of_one_user_is_not_readable_by_another_whatever_the_format(client: AsyncClient, workspace, small_threshold):
    form = await publish(workspace)
    await submit_many(workspace, form, 3)
    export = await queued(client, workspace, form, {"format": "json"})
    assert await build(export["id"]) == "ready"
    other = await client.get(url(workspace, form, f"/{export['id']}"), cookies=_cookies(invited_user))
    assert other.status_code in (403, 404)
    assert ExportStatus.READY == (await container.response_export_repo().get_by_id(export["id"])).status
