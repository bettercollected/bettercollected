"""The XLSX and JSON writers (``xlsx.py``, ``json_export.py``), without a database:
the golden form and responses of ``test_table.py`` go through them and come back
through an independent reader (openpyxl for the workbook, ``json`` for the
document). Everything here is synthetic."""

import datetime as dt
import io
import json
import re
import zipfile
from typing import Any, List

import pytest
from openpyxl import load_workbook

from backend.app.services.response_export import xlsx as xlsx_module
from backend.app.services.response_export.formats import (
    MEDIA_TYPES,
    media_type,
    stored_format,
    write_file,
)
from backend.app.services.response_export.table import (
    ExportResponse,
    build_plan,
    formula_like,
    formula_safe_cell,
)
from backend.app.services.response_export.writer import RowCount
from backend.app.services.response_export.xlsx import (
    SheetFull,
    TooManyColumns,
    column_letters,
    date_serial,
    sheet_name,
    write_xlsx,
    xml_text,
)

from tests.app.response_export.test_table import GIVEN, responses, source_of


async def collect(stored: str, plan, form=None, items=None, counter=None) -> bytes:
    out = b""
    form = form or GIVEN["form"]
    async for piece in write_file(stored, plan, form, source_of(items if items is not None else responses()), counter):
        out += piece
    return out


def workbook_of(data: bytes):
    return load_workbook(io.BytesIO(data), read_only=False)


def one_response(answers=None, **kwargs) -> ExportResponse:
    return ExportResponse(
        response_id=kwargs.pop("response_id", "r1"),
        responder=kwargs.pop("responder", "someone@example.org"),
        created_at=kwargs.pop("created_at", dt.datetime(2026, 10, 1, 10, 0, 5, tzinfo=dt.timezone.utc)),
        answers=answers or {},
        internal_answers=kwargs.pop("internal_answers", {}),
        hidden_fields=kwargs.pop("hidden_fields", {}),
        form_version=kwargs.pop("form_version", 3),
    )


def simple_form(fields: List[dict], hidden=None) -> dict:
    return {
        "title": "Simple",
        "form_id": "f1",
        "builder_version": "v2",
        "hidden_fields": hidden or [],
        "fields": [{"id": "s1", "type": "slide", "properties": {"fields": fields}}],
    }


def text(value: str) -> dict:
    return {"type": "text", "text": value}


# ------------------------------------------------------------------- helpers


def test_column_letters_follow_the_spreadsheet_alphabet():
    assert [column_letters(i) for i in (0, 25, 26, 51, 52, 701, 702, 16383)] == [
        "A", "Z", "AA", "AZ", "BA", "ZZ", "AAA", "XFD",
    ]


def test_xml_text_keeps_the_text_and_drops_what_xml_cannot_carry():
    assert xml_text("a & b < c > d") == "a &amp; b &lt; c &gt; d"
    assert xml_text("a\x00b\x0bc\x1fd\ud800e￾") == "abcde"
    # Excel decodes ``_x0041_`` inside a string: it is escaped so the text survives
    assert xml_text("_x0041_") == "_x005F_x0041_"
    assert xml_text("x_y_x41_") == "x_y_x41_"
    # overlapping escapes (the closing underscore of one opens the next), and a capital X
    assert xml_text("_x0041_x0042_") == "_x005F_x0041_x005F_x0042_"
    assert xml_text("_X0041_") == "_x005F_X0041_"
    assert xml_text("_x004") == "_x004"
    assert len(xml_text("a" * 40_000)) == xlsx_module.MAX_CELL_CHARS


def test_sheet_names_follow_excels_rules():
    taken: List[str] = []
    for raw in ["Jobs", "jobs", "A/B:C*D?[E]\\", "x" * 50, "", "'quoted'", "History", "History"]:
        taken.append(sheet_name(raw, taken))
    assert taken[0] == "Jobs" and taken[1] == "jobs (2)"
    assert taken[2] == "A B C D E"
    assert len(taken[3]) == 31
    assert taken[4] == "Sheet"
    assert taken[5] == "quoted"
    assert taken[6] == "History (2)"
    assert len({t.lower() for t in taken}) == len(taken)
    assert all(len(t) <= 31 for t in taken)
    # Excel counts UTF-16 units: an emoji is two
    emoji = sheet_name("\U0001f600" * 20, [])
    assert emoji == "\U0001f600" * 15 and len(emoji.encode("utf-16-le")) // 2 == 30
    assert len(sheet_name("\U0001f600" * 20, [emoji]).encode("utf-16-le")) // 2 <= 31


def test_a_name_left_equal_once_illegal_characters_are_gone_is_still_unique():
    taken: List[str] = []
    for raw in ["A\ufffe", "A", "a\uffff", "\x00"]:
        taken.append(sheet_name(raw, taken))
    assert taken == ["A", "A (2)", "a (3)", "Sheet"]


def test_illegal_characters_go_before_the_formula_check():
    for value in ("\x00=1+1", "\ufffe@SUM(A1)", "\x01\x02-2+3"):
        xml = xlsx_module.cell_xml("A1", value)
        assert f's="{xlsx_module.STYLE_TEXT}"' in xml, value
        assert "\x00" not in xml and "\ufffe" not in xml
    assert xlsx_module.cell_xml("A1", "\x00\x01") == ""


def test_dates_are_serials_only_when_they_are_iso_dates():
    assert date_serial("2026-10-01", "date") == 46296.0
    assert date_serial("2026-10-01T12:00:00Z", "datetime") == 46296.5
    # a zone is converted to UTC
    assert date_serial("2026-10-01T12:00:00+02:00", "datetime") == pytest.approx(46296 + 10 / 24)
    for bad in ["", "1/10/2026", "2026-13-40", "2026-10-01 extra", "1850-01-01", "tomorrow"]:
        assert date_serial(bad, "date") is None
    assert date_serial("2026-10-01", "datetime") is None


def test_formula_like_matches_what_the_csv_guard_prefixes():
    for risky in ["=1+1", "+31 6", "-x", "@sum", "  =1", "\tx", "\rx", " =1"]:
        assert formula_like(risky) and formula_safe_cell(risky) == "'" + risky
    for safe in ["-5", "-5.25", "plain", "a=b", "", 7]:
        assert not formula_like(safe) and formula_safe_cell(safe) == safe


# ---------------------------------------------------------------------- XLSX


@pytest.mark.asyncio
async def test_a_workbook_has_the_csv_tables_as_sheets_and_the_same_cells():
    plan = build_plan(GIVEN["form"])
    data = await collect("xlsx", plan)
    assert data.startswith(b"PK")
    book = workbook_of(data)
    assert book.sheetnames == xlsx_module.sheet_names(plan)
    assert book.sheetnames[0] == "Responses" and len(book.sheetnames) == 1 + len(plan.group_tables)

    main = list(book["Responses"].iter_rows(values_only=True))
    assert list(main[0]) == plan.headers
    expected = [plan.row(r) for r in responses()]
    assert len(main) == 1 + len(expected)
    kinds = plan.kinds
    for got, want in zip(main[1:], expected):
        assert len(got) == len(want)
        for value, original, kind, header in zip(got, want, kinds, plan.headers):
            if original in ("", None):
                assert value in (None, "")
            elif kind == "datetime":
                assert isinstance(value, dt.datetime)
                assert value.strftime("%Y-%m-%dT%H:%M:%SZ") == original
            elif kind == "date" and isinstance(value, dt.datetime):
                assert value.strftime("%Y-%m-%d") == original
            elif isinstance(original, (int, float)):
                assert value == original, header
            else:
                # text is stored as the respondent wrote it: no apostrophe, nothing evaluated
                assert str(value) == str(original), header

    # each repeating group exported as a table is a sheet of its own, like the CSV archive
    for sheet, table in zip(book.sheetnames[1:], plan.group_tables):
        rows = list(book[sheet].iter_rows(values_only=True))
        assert list(rows[0]) == table.headers
        assert len(rows) - 1 == sum(len(table.rows(r)) for r in responses())


@pytest.mark.asyncio
async def test_the_submission_time_and_a_date_answer_are_real_dates():
    plan = build_plan(GIVEN["form"])
    sheet = workbook_of(await collect("xlsx", plan))["Responses"]
    headers = [c.value for c in sheet[1]]
    submitted = headers.index(next(h for h in headers if h.startswith("Submitted at")))
    first = sheet[2][submitted]
    assert first.is_date and first.value == dt.datetime(2026, 10, 1, 10, 0, 5)
    assert first.number_format == "yyyy\\-mm\\-dd\\ hh:mm:ss"
    date_header = headers.index("Date")
    cell = next((row[date_header] for row in sheet.iter_rows(min_row=2) if row[date_header].value), None)
    assert cell is not None and cell.is_date and cell.number_format == "yyyy\\-mm\\-dd"


@pytest.mark.asyncio
async def test_text_is_never_a_formula_and_a_risky_cell_carries_the_text_flag():
    risky = ["=HYPERLINK(\"http://x\",\"y\")", "+31 6 1234 5678", "-1+1", "@SUM(A1)", "  =1+1", "\t=1", "-5", "fine"]
    form = simple_form([{"id": f"q{i}", "type": "short_text", "title": f"Q{i}"} for i in range(len(risky))])
    plan = build_plan(form)
    response = one_response({f"q{i}": text(v) for i, v in enumerate(risky)})
    book = workbook_of(await collect("xlsx", plan, form, [response]))
    row = {c.column_letter: c for c in book["Responses"][2]}
    # no formula anywhere; the value is intact; only a formula-like string is flagged as text
    for i, original in enumerate(risky):
        cell = row[xlsx_module.column_letters(i + 1)]
        assert cell.data_type == "s" and cell.value == original
        assert bool(cell.quotePrefix) == formula_like(original)
    sheet_xml = zipfile.ZipFile(io.BytesIO(await collect("xlsx", plan, form, [response]))).read("xl/worksheets/sheet1.xml").decode()
    assert "<f>" not in sheet_xml and "<f " not in sheet_xml


@pytest.mark.asyncio
async def test_a_header_that_looks_like_a_formula_is_text_too():
    form = simple_form([{"id": "q", "type": "short_text", "title": "=1+1"}])
    book = workbook_of(await collect("xlsx", build_plan(form), form, [one_response({"q": text("x")})]))
    header = book["Responses"]["B1"]
    assert header.value == "=1+1" and header.data_type == "s" and header.quotePrefix and header.font.b


@pytest.mark.asyncio
async def test_awkward_text_survives_the_round_trip():
    awkward = [
        "plain",
        "tab\there\nnew line",
        "  padded  ",
        "<b>markup</b> & more",
        "control\x00\x01\x1f gone",
        "emoji \U0001f600 and accents éè",
        "lone \ud800 surrogate",
        "x" * 40_000,
    ]
    form = simple_form([{"id": f"q{i}", "type": "long_text", "title": f"Q{i}"} for i in range(len(awkward))])
    book = workbook_of(await collect("xlsx", build_plan(form), form, [one_response({f"q{i}": text(v) for i, v in enumerate(awkward)})]))
    values = [c.value for c in book["Responses"][2]][1:-2]
    assert values[0] == "plain"
    assert values[1] == "tab\there\nnew line"
    assert values[2] == "  padded  "
    assert values[3] == "<b>markup</b> & more"
    assert values[4] == "control gone"
    assert values[5] == "emoji \U0001f600 and accents éè"
    assert values[6] == "lone  surrogate"
    assert len(values[7]) == xlsx_module.MAX_CELL_CHARS


@pytest.mark.asyncio
async def test_the_escape_excel_decodes_in_a_string_is_itself_escaped():
    # Excel and LibreOffice read ``_x0041_`` in a string as "A" (ECMA-376 ST_Xstring); the
    # text must reach the person as written, so the underscore is escaped. (openpyxl
    # does not decode the escape in an inline string, so the part is read raw here; the
    # build was also checked by converting a file with LibreOffice.)
    form = simple_form([{"id": "q", "type": "short_text", "title": "Q"}])
    data = await collect("xlsx", build_plan(form), form, [one_response({"q": text("_x0041_ and _x005F_ and _x")})])
    xml = zipfile.ZipFile(io.BytesIO(data)).read("xl/worksheets/sheet1.xml").decode()
    assert "_x005F_x0041_ and _x005F_x005F_ and _x" in xml


@pytest.mark.asyncio
async def test_a_carriage_return_and_a_hostile_sheet_name_survive():
    form = simple_form(
        [
            {"id": "q", "type": "long_text", "title": "Q"},
            {
                "id": "g",
                "type": "group",
                "title": 'Say "hi" & <b>\'x\'</b> [a]',
                "properties": {
                    "repeat": {"max_items": 6, "item_label": "Item", "export_layout": "rows"},
                    "fields": [{"id": "c", "type": "short_text", "title": "C"}],
                },
            },
        ]
    )
    given = {"q": text("one\r\ntwo\rthree"), "g": {"type": "group", "items": [{"c": text("x")}]}}
    data = await collect("xlsx", build_plan(form), form, [one_response(given)])
    book = workbook_of(data)
    # [ ] and / are not allowed in a sheet name; a quote, & and < are, and must not break the file
    assert book.sheetnames[0] == "Responses" and len(book.sheetnames) == 2
    assert book.sheetnames[1].startswith('Say "hi" & <b>\'x\'< b>')
    assert book["Responses"]["B2"].value == "one\r\ntwo\rthree"
    assert [c.value for c in book[book.sheetnames[1]][2]][1:] == [1, "x"]


@pytest.mark.asyncio
async def test_numbers_stay_numbers_and_long_integers_become_text():
    form = simple_form([{"id": "n", "type": "number", "title": "N"}, {"id": "m", "type": "number", "title": "M"}])
    plan = build_plan(form)
    data = await collect("xlsx", plan, form, [one_response({"n": {"type": "number", "number": 12}, "m": {"type": "number", "number": 12345678901234567}})])
    row = workbook_of(data)["Responses"][2]
    assert row[1].value == 12 and row[1].data_type == "n"
    # Excel keeps 15 digits: a longer integer is written as text (read raw: openpyxl guesses types)
    xml = zipfile.ZipFile(io.BytesIO(data)).read("xl/worksheets/sheet1.xml").decode()
    assert '<c r="B2"><v>12</v></c>' in xml
    assert 't="inlineStr"><is><t xml:space="preserve">12345678901234567</t>' in xml


@pytest.mark.asyncio
async def test_the_workbook_is_streamed_in_pieces_and_counts_the_main_table():
    plan = build_plan(GIVEN["form"])
    many = [one_response({"name": text(f"n{i}")}, response_id=f"r{i}") for i in range(400)]
    counter = RowCount()
    pieces = []
    async for piece in write_xlsx(plan, "T", source_of(many), counter, piece_bytes=2048):
        pieces.append(piece)
    assert counter.rows == 400
    assert len(pieces) > 5
    assert max(len(p) for p in pieces) < 64 * 1024
    assert len(list(workbook_of(b"".join(pieces))["Responses"].iter_rows())) == 401


@pytest.mark.asyncio
async def test_the_workbook_holds_no_formula_name_link_or_macro_part():
    data = await collect("xlsx", build_plan(GIVEN["form"]))
    archive = zipfile.ZipFile(io.BytesIO(data))
    assert archive.testzip() is None
    names = archive.namelist()
    assert names[:5] == [
        "[Content_Types].xml", "_rels/.rels", "xl/workbook.xml", "xl/_rels/workbook.xml.rels", "xl/styles.xml",
    ]
    assert not any("vba" in n.lower() or "external" in n.lower() or n.endswith(".bin") for n in names)
    workbook = archive.read("xl/workbook.xml").decode()
    assert "definedName" not in workbook and "externalReference" not in workbook
    for name in names:
        if name.startswith("xl/worksheets/"):
            xml = archive.read(name).decode()
            # no formula element, and no hyperlink part or element: text only
            assert not re.search(r"<f[ >/]", xml) and "<hyperlink" not in xml


@pytest.mark.asyncio
async def test_a_sheet_past_excels_row_limit_stops_the_file(monkeypatch):
    monkeypatch.setattr(xlsx_module, "MAX_ROWS_PER_SHEET", 3)
    plan = build_plan(GIVEN["form"])
    with pytest.raises(SheetFull):
        await collect("xlsx", plan, items=[one_response() for _ in range(5)])


@pytest.mark.asyncio
async def test_more_columns_than_a_sheet_holds_is_refused_before_a_byte(monkeypatch):
    monkeypatch.setattr(xlsx_module, "MAX_COLUMNS", 4)
    plan = build_plan(GIVEN["form"])
    with pytest.raises(TooManyColumns):
        async for _ in write_xlsx(plan, "T", source_of(responses())):
            raise AssertionError("no byte may be written")


@pytest.mark.asyncio
async def test_chosen_columns_are_the_only_columns_of_the_sheet():
    plan = build_plan(GIVEN["form"], ["responder", "q:name", "response_id"])
    book = workbook_of(await collect("xlsx", plan))
    assert [c.value for c in book["Responses"][1]] == plan.headers and len(plan.headers) == 3
    assert book.sheetnames == ["Responses"]


# ---------------------------------------------------------------------- JSON


async def json_of(stored: str, plan, form=None, items=None) -> Any:
    return json.loads((await collect(stored, plan, form, items)).decode("utf-8"))


@pytest.mark.asyncio
async def test_the_json_document_follows_the_apis_response_shape():
    plan = build_plan(GIVEN["form"])
    document = await json_of("json", plan)
    assert set(document) == {"form", "questions", "responses"}
    assert document["form"]["title"] == GIVEN["form"]["title"]
    assert len(document["responses"]) == len(GIVEN["responses"])
    first = document["responses"][0]
    assert first["responseId"] == "resp-1" and first["submittedAt"] == "2026-10-01T10:00:05Z"
    assert first["responder"] == "sita@example.org"
    # the stored answer objects, keyed by question id (what the REST API returns)
    assert first["answers"]["phone"] == {"type": "phone_number", "phone_number": "+31 6 1234 5678"}
    assert first["answers"]["colours"]["choices"]["values"] == ["c1", "c2", "c3"]
    assert first["hiddenFields"]["ref"] == "=SUM(1)"  # JSON has no formula problem
    # the repeating group stays nested
    kids = first["answers"]["kids"]["items"]
    assert isinstance(kids, list) and all(isinstance(item, dict) for item in kids)


@pytest.mark.asyncio
async def test_the_questions_say_what_the_ids_are():
    document = await json_of("json", build_plan(GIVEN["form"]))
    by_id = {q["id"]: q for q in document["questions"]}
    assert by_id["colours"]["type"] == "multiple_choice"
    assert {c["id"] for c in by_id["colours"]["choices"]} >= {"c1", "c2", "c3"}
    assert all(set(c) == {"id", "label"} for c in by_id["colours"]["choices"])
    assert [f["id"] for f in by_id["kids"]["fields"]] == ["kname", "kage", "kcol"]
    assert by_id["ref"].get("internal") is True


@pytest.mark.asyncio
async def test_json_holds_only_the_chosen_columns():
    plan = build_plan(GIVEN["form"], ["q:name", "hidden:utm_source"])
    document = await json_of("json", plan)
    assert [q["id"] for q in document["questions"]] == ["name"]
    for item in document["responses"]:
        assert set(item) <= {"responseId", "submittedAt", "formVersion", "answers", "hiddenFields"}
        assert "responder" not in item and "internalAnswers" not in item
        assert set(item["answers"]) <= {"name"}
        assert set(item["hiddenFields"]) <= {"utm_source"}


@pytest.mark.asyncio
async def test_json_never_carries_a_link_to_a_stored_file():
    form = simple_form(
        [
            {"id": "sig", "type": "signature", "title": "Sign"},
            {"id": "up", "type": "file_upload", "title": "Upload"},
        ]
    )
    answers = {
        "sig": {"type": "signature", "signature": {"method": "drawn", "file_id": "abc123", "url": "/workspaces/w/submissions/r/signatures/sig", "signed_at": "2026-10-01T10:00:00Z"}},
        "up": {
            "type": "file_url",
            "file_url": "https://storage.example/obj?sig=secret",
            "file_metadata": {"id": "f9", "name": "cv.pdf", "size": 10, "url": "https://storage.example/obj?sig=secret"},
        },
    }
    raw = (await collect("json", build_plan(form), form, [one_response(answers)])).decode()
    assert "https://" not in raw and "secret" not in raw and "abc123" not in raw and "/signatures/" not in raw
    document = json.loads(raw)
    answer = document["responses"][0]["answers"]
    assert answer["sig"]["signature"] == {"method": "drawn", "signed_at": "2026-10-01T10:00:00Z"}
    assert answer["up"]["file_metadata"] == {"id": "f9", "name": "cv.pdf", "size": 10}


@pytest.mark.asyncio
async def test_json_is_valid_utf8_whatever_the_answers_hold():
    form = simple_form([{"id": "q", "type": "short_text", "title": "Lone \ud800 title"}])
    raw = await collect("json", build_plan(form), form, [one_response({"q": text("lone \udc00 é \U0001f600")}, responder="r\ud800@x")])
    document = json.loads(raw.decode("utf-8"))
    item = document["responses"][0]
    assert item["answers"]["q"]["text"] == "lone � é \U0001f600"
    assert item["responder"] == "r�@x"
    assert "NaN" not in raw.decode() and "Infinity" not in raw.decode()


@pytest.mark.asyncio
async def test_json_of_no_responses_is_an_empty_list_and_counts_rows():
    counter = RowCount()
    out = b""
    async for piece in write_file("json", build_plan(GIVEN["form"]), GIVEN["form"], source_of([]), counter):
        out += piece
    assert json.loads(out)["responses"] == [] and counter.rows == 0
    counter = RowCount()
    async for _ in write_file("json", build_plan(GIVEN["form"]), GIVEN["form"], source_of(responses()), counter):
        pass
    assert counter.rows == len(GIVEN["responses"])


# --------------------------------------------------------------- the formats


def test_the_stored_format_and_media_type_follow_the_request():
    with_tables = build_plan(GIVEN["form"])
    assert with_tables.group_tables
    assert stored_format("csv", with_tables) == "zip" and stored_format(None, with_tables) == "zip"
    assert stored_format("xlsx", with_tables) == "xlsx" and stored_format("json", with_tables) == "json"
    plain = build_plan(simple_form([{"id": "q", "type": "short_text", "title": "Q"}]))
    assert stored_format("csv", plain) == "csv"
    assert media_type("xlsx") == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert media_type("json").startswith("application/json")
    assert media_type("zip") == "application/zip" and media_type("csv").startswith("text/csv")
    assert media_type("unknown") == MEDIA_TYPES["csv"]
